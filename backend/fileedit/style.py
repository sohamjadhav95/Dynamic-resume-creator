"""Measure the design of an uploaded resume so the app can draw it in the same style.

Works on any single-column resume. Returns sizes and distances in PDF points in the same shape the
web page engine uses (web/src/lib/page/layout.ts, interface Style). Anything that cannot be
measured is left out, and the engine uses its default for it.
"""

import io
import logging
import re
import statistics
from collections import Counter

import docx
import pymupdf

from .common import CONTACT_RE, DATE_RE, clean, looks_like_heading
from .pdf_edit import MONO_RE, build_groups

log = logging.getLogger("resume.style")

SERIF_RE = re.compile(
    r"times|georgia|garamond|cambria|book|serif|minion|palatino|caslon|baskerville|didot|merriweather|lora|roman|"
    r"cmr|lmr|charter|charis|bodoni|constantia|crimson|playfair|libre ?baskerville|pt ?serif|source ?serif|"
    r"noto ?serif|droid ?serif|georgia|tinos|liberation ?serif|dejavu ?serif|century|bookman|goudy|sabon|utopia",
    re.I,
)

SANS_RE = re.compile(
    r"sans|arial|helvetica|calibri|verdana|tahoma|roboto|inter|lato|segoe|gill|futura|avenir|montserrat|"
    r"poppins|nunito|raleway|ubuntu|carlito|arimo|aptos|franklin|trebuchet|century gothic|frutiger|myriad",
    re.I,
)

# Which section a heading text belongs to (so the resume keeps its own heading words).
SECTION_RE = [
    ("summary", re.compile(r"summary|profile|about|objective", re.I)),
    ("skills", re.compile(r"skill|competenc|technolog|tools|expertise|tech stack", re.I)),
    ("experience", re.compile(r"experience|employment|work history|career", re.I)),
    ("projects", re.compile(r"project", re.I)),
    ("certifications", re.compile(r"certif|licen|achievement|award|course", re.I)),
    ("education", re.compile(r"education|academic|qualification", re.I)),
]


def _clamp(v: float, lo: float, hi: float) -> float:
    return round(max(lo, min(hi, v)), 2)


def _hex(color: int) -> str:
    return f"#{color:06x}"


def _median(values: list[float], default: float) -> float:
    return statistics.median(values) if values else default


def _section_titles(headings: list[str]) -> dict[str, str]:
    titles: dict[str, str] = {}
    for text in headings:
        for key, rx in SECTION_RE:
            if rx.search(text) and key not in titles:
                titles[key] = clean(text).rstrip(":")
                break
    return titles


# ------------------------------------------------------------------ PDF


def detect_pdf_style(data: bytes) -> dict:
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        page = doc[0]
        W, H = page.rect.width, page.rect.height
        groups, body = build_groups(doc)
        groups = [g for g in groups if g.page == 0]
        lines = [ln for g in groups for ln in g.lines]
        if len(lines) < 5:
            return {}

        # Fonts: serif or sans, from the spans that carry most of the text.
        families: Counter = Counter()
        span_info = page.get_text("dict")["blocks"]
        for b in span_info:
            for ln in b.get("lines", []):
                for sp in ln["spans"]:
                    n = len(sp["text"].strip())
                    if not n:
                        continue
                    name = sp["font"]
                    mono = bool(MONO_RE.search(name)) or (bool(sp["flags"] & 8) and not SANS_RE.search(name))
                    # The font's name is more reliable than its flags (some sans fonts claim to be serif).
                    serif = bool(SERIF_RE.search(name)) or (not SANS_RE.search(name) and bool(sp["flags"] & 4))
                    families["mono" if mono else "serif" if serif else "sans"] += n
        family = families.most_common(1)[0][0] if families else "sans"
        if family == "mono":
            family = "sans"

        style: dict = {"family": family, "page": "Letter" if abs(W - 612) < 10 and abs(H - 792) < 10 else "A4"}

        # Margins from where the text sits.
        lefts = sorted(ln.x0 for ln in lines)
        rights = sorted(ln.x1 for ln in lines)
        left = lefts[max(0, len(lefts) // 20)]
        right = W - rights[min(len(rights) - 1, len(rights) * 19 // 20)]
        top = min(ln.y0 for ln in lines)
        bottom_text = max(ln.y1 for ln in lines)
        # Resumes are almost always symmetric; the left edge is reliable (ragged-right text is not).
        side = left if right > left + 6 else (left + right) / 2
        top_m = top - 2
        # A page that is not full says nothing about its bottom margin: assume it matches the top.
        bottom_m = H - bottom_text if H - bottom_text < top_m else max(top_m, 24)
        style["margin"] = {
            "side": _clamp(side, 18, 90),
            "top": _clamp(top_m, 14, 90),
            "bottom": _clamp(bottom_m, 18, 60),
        }

        def size_of(g) -> float:
            sizes = Counter(round(c.size * 2) / 2 for ln in g.lines for c in ln.chars if c.c.strip())
            return sizes.most_common(1)[0][0] if sizes else body

        def is_bold(g) -> bool:
            chars = [c for ln in g.lines for c in ln.chars if c.c.strip()]
            return bool(chars) and sum(c.bold for c in chars) > 0.6 * len(chars)

        def is_italic(g) -> bool:
            chars = [c for ln in g.lines for c in ln.chars if c.c.strip()]
            return bool(chars) and sum(c.italic for c in chars) > 0.6 * len(chars)

        def color_of(g) -> int:
            cols = Counter(c.color for ln in g.lines for c in ln.chars if c.c.strip())
            return cols.most_common(1)[0][0] if cols else 0

        def baseline(g, last=False) -> float:
            return g.lines[-1 if last else 0].baseline

        # Header: the name is the biggest text near the top, then contact / tagline lines.
        top_groups = [g for g in groups[:8] if g.lines[0].y0 < H * 0.3]
        name_g = max(top_groups or groups[:3], key=size_of)
        ni = groups.index(name_g)
        style["name"] = _clamp(size_of(name_g), 12, 36)
        centre = (name_g.lines[0].x0 + name_g.lines[0].x1) / 2
        style["headerAlign"] = "center" if abs(centre - W / 2) < W * 0.08 else "left"
        heading_idx = [i for i, g in enumerate(groups) if g.heading and i > ni]
        first_heading = heading_idx[0] if heading_idx else len(groups)
        header = groups[ni + 1 : first_heading]
        contact = next((g for g in header if CONTACT_RE.search(g.full_text)), None)
        tagline = next((g for g in header if g is not contact and not CONTACT_RE.search(g.full_text)), None)
        if contact:
            style["contact"] = _clamp(size_of(contact), 6, 12)
        if tagline:
            style["badge"] = _clamp(size_of(tagline), 7, 14)
            style["badgeBold"] = is_bold(tagline)
            style["badgeItalic"] = is_italic(tagline)

        # Section headings: size, case, colour, and the rule line under them.
        heads = [groups[i] for i in heading_idx]
        if heads:
            style["title"] = _clamp(_median([size_of(g) for g in heads], body + 1), 7, 18)
            style["titleBold"] = sum(is_bold(g) for g in heads) >= len(heads) / 2
            colour = Counter(color_of(g) for g in heads).most_common(1)[0][0]
            style["accent"] = _hex(colour)
            style["titleUpper"] = False  # the headings' own words are kept, with their own case
            style["titles"] = _section_titles([g.full_text for g in heads])
            rules = []
            for d in page.get_drawings():
                for item in d.get("items", []):
                    kind = item[0]
                    if kind == "l":
                        p1, p2 = item[1], item[2]
                        if abs(p1.y - p2.y) < 1 and abs(p2.x - p1.x) > W * 0.4:
                            rules.append((p1.y, d.get("width") or 0.75, d.get("color")))
                    elif kind == "re":
                        r = item[1]
                        if r.height < 3 and r.width > W * 0.4:
                            rules.append((r.y0 + r.height / 2, max(r.height, 0.4), d.get("fill") or d.get("color")))
            under = []
            for g in heads:
                y1 = g.lines[-1].y1
                hit = [r for r in rules if y1 - 3 <= r[0] <= y1 + 9]
                if hit:
                    under.append(hit[0])
            style["rule"] = (
                _clamp(_median([r[1] for r in under], 0.75), 0.3, 2.5) if len(under) >= len(heads) / 2 else 0
            )

        # Body text, line spacing and bullets.
        style["body"] = _clamp(body, 7, 13)
        pitches = []
        for g in groups:
            if len(g.lines) >= 2 and abs(size_of(g) - body) < 0.6:
                for a, b in zip(g.lines, g.lines[1:], strict=False):
                    pitches.append((b.baseline - a.baseline) / body)
        lh = _clamp(_median(pitches, 1.25), 1.08, 1.6)
        style["lineHeight"] = lh
        bullets = [g for g in groups if g.is_bullet]
        if bullets:
            bx = [g.lines[0].chars[0].x0 for g in bullets]
            tx = [g.body_chars(0)[0].x0 for g in bullets if g.body_chars(0)]
            style["bulletIndent"] = _clamp(_median(bx, left) - left, 0, 30)
            style["textIndent"] = _clamp(_median(tx, left + 14) - left, 6, 40)
            glyph = Counter(clean(g.prefix)[:1] for g in bullets).most_common(1)[0][0]
            style["bullet"] = {"●": "•", "▪": "•", "■": "•", "◦": "•", "": "•", "‣": "•", "·": "•"}.get(
                glyph, glyph or "•"
            )

        # Entry titles (bold lines that are not headings), tech-stack lines (italic) and skills lines.
        bold_rows = [
            g
            for g in groups
            if not g.heading and not g.is_bullet and is_bold(g) and g not in (name_g, contact, tagline)
        ]
        if bold_rows:
            style["heading"] = _clamp(_median([size_of(g) for g in bold_rows], body), 7, 14)
            style["role"] = style["heading"]
        italics = [g for g in groups if not g.heading and is_italic(g) and g is not tagline]
        if italics:
            style["stack"] = _clamp(_median([size_of(g) for g in italics], body), 6.5, 13)
        skills = [g for g in groups if g.kind == "skills"]
        if skills:
            style["skills"] = _clamp(_median([size_of(g) for g in skills], body), 7, 13)

        # Distances, converted to the engine's "gap between line boxes".
        def offset(size: float) -> float:
            pitch = size * lh
            return size * 0.86 + (pitch - size) / 2

        def tail(size: float) -> float:
            return size * lh - offset(size)

        title_size = style.get("title", body + 1)
        section, after = [], []
        for i in heading_idx:
            g = groups[i]
            if i > 0 and groups[i - 1] is not name_g and i - 1 >= first_heading:
                prev = groups[i - 1]
                d = baseline(g) - baseline(prev, last=True)
                section.append(d - tail(size_of(prev)) - offset(title_size))
            if i + 1 < len(groups):
                nxt = groups[i + 1]
                d = baseline(nxt) - baseline(g, last=True)
                rule = style.get("rule", 0)
                after.append(d - tail(title_size) - (0.6 + rule if rule else 0) - offset(size_of(nxt)))
        entry, bullet_gaps = [], []
        for a, b in zip(groups, groups[1:], strict=False):
            d = baseline(b) - baseline(a, last=True) - tail(size_of(a)) - offset(size_of(b))
            if a.is_bullet and b.is_bullet:
                bullet_gaps.append(d)
            elif a.is_bullet and b in bold_rows:
                entry.append(d)
        style["gaps"] = {
            "section": _clamp(_median(section, 8), 2, 24),
            "afterTitle": _clamp(_median(after, 3), 0, 14),
            "entry": _clamp(_median(entry, 5), 0, 16),
            "bullet": _clamp(_median(bullet_gaps, 2), 0, 8),
        }
        if heading_idx:
            last_header = groups[first_heading - 1]
            d = baseline(groups[first_heading]) - baseline(last_header, last=True)
            style["gaps"]["header"] = _clamp(d - tail(size_of(last_header)) - offset(title_size), 2, 30)
        if contact:
            d = baseline(contact) - baseline(name_g, last=True)
            style["gaps"]["name"] = _clamp(d - tail(style["name"]) - offset(size_of(contact)), 0, 14)
        if contact and tagline and header.index(tagline) > header.index(contact):
            d = baseline(tagline) - baseline(contact, last=True)
            style["gaps"]["contact"] = _clamp(d - tail(size_of(contact)) - offset(size_of(tagline)), 0, 14)
        style["ink"] = _hex(Counter(color_of(g) for g in groups if not g.heading).most_common(1)[0][0])
        return style
    finally:
        doc.close()


# ------------------------------------------------------------------ Word


def detect_docx_style(data: bytes) -> dict:
    d = docx.Document(io.BytesIO(data))
    sec = d.sections[0]
    pt = 12700  # EMU per point
    style: dict = {}
    width, height = sec.page_width / pt, sec.page_height / pt
    style["page"] = "Letter" if abs(width - 612) < 10 and abs(height - 792) < 10 else "A4"
    style["margin"] = {
        "side": _clamp((sec.left_margin + sec.right_margin) / 2 / pt, 18, 90),
        "top": _clamp(sec.top_margin / pt, 14, 90),
        "bottom": _clamp(sec.bottom_margin / pt, 18, 60),
    }
    normal = d.styles["Normal"]
    font = normal.font.name or "Calibri"
    style["family"] = "serif" if SERIF_RE.search(font) and not SANS_RE.search(font) else "sans"

    from .docx_edit import read_paras  # local import: avoids a cycle with the editing helpers

    paras, body = read_paras(d)
    if not paras:
        return style
    style["body"] = _clamp(body, 7, 13)
    spacing = normal.paragraph_format.line_spacing
    multiple = spacing if isinstance(spacing, float) and spacing < 4 else 1.0
    style["lineHeight"] = _clamp(1.17 * multiple, 1.08, 1.6)
    families = Counter(
        ("serif" if SERIF_RE.search(p.font or "") and not SANS_RE.search(p.font or "") else "sans")
        for p in paras
        for _ in p.text
    )
    style["family"] = families.most_common(1)[0][0]

    name = max(paras[:4], key=lambda p: p.size)
    style["name"] = _clamp(name.size, 12, 36)
    align = d.paragraphs[0].alignment if d.paragraphs else None
    style["headerAlign"] = "center" if align is not None and int(align) == 1 else "left"
    heads = [p for p in paras if p is not name and not p.is_bullet and looks_like_heading(p.text, p.bold, p.size, body)]
    if heads:
        style["title"] = _clamp(_median([p.size for p in heads], body + 1), 7, 18)
        style["titleUpper"] = False
        style["titles"] = _section_titles([p.text for p in heads])
        style["rule"] = (
            0.75
            if any(
                p.el.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr") is not None
                for p in heads
            )
            else 0
        )
    contact = next((p for p in paras[1:6] if CONTACT_RE.search(p.text)), None)
    if contact:
        style["contact"] = _clamp(contact.size, 6, 12)
    rows = [
        p for p in paras if p.bold and not p.is_bullet and p not in heads and p is not name and DATE_RE.search(p.text)
    ]
    if rows:
        style["heading"] = _clamp(_median([p.size for p in rows], body), 7, 14)
        style["role"] = style["heading"]
    return style


def detect_style(file_type: str, data: bytes) -> dict:
    """Never fails: an unusual file just gets the default style."""
    try:
        return detect_pdf_style(data) if file_type == "pdf" else detect_docx_style(data)
    except Exception as exc:  # odd files: fall back to the default look
        log.warning("style.detect_failed type=%s error=%r", file_type, exc)
        return {}
