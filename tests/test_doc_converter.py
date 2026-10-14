import base64
import io

import docx
from fastapi.testclient import TestClient

from backend.app import app

client = TestClient(app)


def test_convert_doc_docx():
    doc = docx.Document()
    doc.add_paragraph("Soham Jadhav")
    doc.add_paragraph("Machine Learning Engineer")
    out = io.BytesIO()
    doc.save(out)
    docx_bytes = out.getvalue()

    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    response = client.post(
        "/api/convert-doc",
        files={"file": ("resume.docx", docx_bytes, mime)},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["ok"] is True
    assert "Soham Jadhav" in data["text"]
    assert "Machine Learning Engineer" in data["text"]
    assert "docx_base64" in data


def test_tailor_doc_bullets():
    doc = docx.Document()
    doc.add_paragraph("Old bullet about machine learning")
    out = io.BytesIO()
    doc.save(out)
    b64_in = base64.b64encode(out.getvalue()).decode("utf-8")

    response = client.post(
        "/api/tailor-doc",
        json={
            "docx_base64": b64_in,
            "replacements": [
                {
                    "old_text": "Old bullet about machine learning",
                    "new_text": "Tailored bullet for Generative AI systems",
                }
            ],
        },
    )
    assert response.status_code == 200
    data = response.json()
    assert data["ok"] is True
    out_bytes = base64.b64decode(data["docx_base64"])
    doc2 = docx.Document(io.BytesIO(out_bytes))
    assert "Tailored bullet for Generative AI systems" in doc2.paragraphs[0].text
