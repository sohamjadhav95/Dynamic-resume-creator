"""In-place editing of uploaded PDF / Word resumes: layout must never change."""

import base64
import io
import json

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from backend import app as app_module
from backend import fileedit, limits
from backend.ai.providers import Chunk
from backend.config import get_settings

BULLET_1 = (
    "Built a data pipeline in Python and Airflow that loads sales data from 12 stores every night, cutting report "
    "time by 40% and removing manual spreadsheet work for the finance team."
)
BULLET_2 = "Set up Docker based test servers for 3 product teams."
SKILLS = "Python, SQL, Airflow, Docker, Tableau"
JD = "We are hiring a Data Engineer with Python, Airflow, SQL and Docker experience to build pipelines. " * 2


def _wrap(text: str, font: pymupdf.Font, size: float, width: float) -> list[str]:
    lines, cur = [], ""
    for word in text.split():
        trial = f"{cur} {word}".strip()
        if cur and font.text_length(trial, size) > width:
            lines.append(cur)
            cur = word
        else:
            cur = trial
    return lines + [cur]


def make_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    helv = pymupdf.Font("helv")
    page.insert_text((230, 50), "Jane Doe", fontname="hebo", fontsize=18)
    page.insert_text((50, 70), "jane@example.com | github.com/jane", fontname="helv", fontsize=9)
    page.insert_text((50, 100), "EXPERIENCE", fontname="hebo", fontsize=11)
    page.insert_text((50, 116), "Data Engineer | Acme Retail", fontname="hebo", fontsize=10)
    y = 132.0
    for bullet in (BULLET_1, BULLET_2):
        lines = _wrap(bullet, helv, 10, 485)
        page.insert_text((50, y), "•", fontname="helv", fontsize=10)
        for line in lines:
            page.insert_text((60, y), line, fontname="helv", fontsize=10)
            y += 12
        y += 4
    page.insert_text((50, y + 14), "SKILLS", fontname="hebo", fontsize=11)
    label = "Languages: "
    page.insert_text((50, y + 30), label, fontname="hebo", fontsize=10)
    page.insert_text((50 + pymupdf.Font("hebo").text_length(label, 10), y + 30), SKILLS, fontname="helv", fontsize=10)
    return doc.tobytes()


def make_docx() -> bytes:
    d = docx.Document()
    d.add_paragraph().add_run("Jane Doe").bold = True
    d.add_paragraph().add_run("EXPERIENCE").bold = True
    p = d.add_paragraph()
    p.add_run("•")
    p.add_run("\t" + BULLET_1)
    d.add_paragraph().add_run("SKILLS").bold = True
    p = d.add_paragraph()
    p.add_run("Languages: ").bold = True
    p.add_run(SKILLS)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def seg(extracted, kind):
    return next(s for s in extracted.segments if s.kind == kind)


def pdf_text(data: bytes) -> str:
    return " ".join(pymupdf.open(stream=data, filetype="pdf")[0].get_text().split())


# ---------------------------------------------------------------- PDF


def test_pdf_finds_bullets_and_skills_but_not_names():
    ex = fileedit.extract("pdf", make_pdf())
    bullet = seg(ex, "bullet")
    assert bullet.lines >= 2 and bullet.text.startswith("Built a data pipeline")
    skills = seg(ex, "skills")
    assert skills.prefix == "Languages:" and skills.text == SKILLS
    editable = " ".join(s.text for s in ex.segments)
    assert "Jane Doe" not in editable and "jane@example.com" not in editable and "Acme" not in editable


def test_pdf_same_size_rewrite_is_written_in_place():
    data = make_pdf()
    bullet = seg(fileedit.extract("pdf", data), "bullet")
    new = (
        "Built a nightly Python and Airflow pipeline that loads sales data from 12 stores, cutting report time by "
        "40% and ending manual spreadsheet work for the finance team members."
    )
    out, results = fileedit.apply("pdf", data, {bullet.id: new})
    assert results[0].status == "applied", results[0].note
    assert pymupdf.open(stream=out, filetype="pdf").page_count == 1
    text = pdf_text(out)
    assert "nightly Python and Airflow pipeline" in text
    assert "removing manual spreadsheet work" not in text
    for kept in ("Jane Doe", "EXPERIENCE", "Data Engineer | Acme Retail", BULLET_2, "Languages:"):
        assert kept in text


def test_pdf_refuses_longer_text_and_dropped_numbers():
    data = make_pdf()
    bullet = seg(fileedit.extract("pdf", data), "bullet")
    too_long = BULLET_1 + " It also sent a short summary email to every store manager each morning."
    _, results = fileedit.apply("pdf", data, {bullet.id: too_long})
    assert results[0].status == "too_long"
    no_number = BULLET_1.replace("by 40%", "a lot")
    _, results = fileedit.apply("pdf", data, {bullet.id: no_number})
    assert results[0].status == "changed_facts" and "40" in results[0].note


def test_pdf_skills_keep_their_label():
    data = make_pdf()
    skills = seg(fileedit.extract("pdf", data), "skills")
    out, results = fileedit.apply("pdf", data, {skills.id: "Python, Airflow, SQL, Docker, Tableau"})
    assert results[0].status == "applied"
    assert "Languages: Python, Airflow, SQL, Docker, Tableau" in pdf_text(out)


def test_pdf_only_editable_lines_can_change():
    data = make_pdf()
    ex = fileedit.extract("pdf", data)
    name_id = next(o["id"] for o in ex.outline if o["text"] == "Jane Doe")
    _, results = fileedit.apply("pdf", data, {name_id: "John Smith"})
    assert results[0].status == "not_allowed"


# ---------------------------------------------------------------- Word


def test_docx_rewrite_keeps_bullet_and_bold_label():
    data = make_docx()
    ex = fileedit.extract("docx", data)
    bullet, skills = seg(ex, "bullet"), seg(ex, "skills")
    new_bullet = (
        "Built a nightly Python and Airflow pipeline that loads sales data from 12 stores, cutting report time by "
        "40% and ending manual spreadsheet work for the finance team."
    )
    out, results = fileedit.apply(
        "docx", data, {bullet.id: new_bullet, skills.id: "Python, Airflow, SQL, Docker, Tableau"}
    )
    assert [r.status for r in results] == ["applied", "applied"]
    paras = [p for p in docx.Document(io.BytesIO(out)).paragraphs if p.text.strip()]
    assert paras[2].text == "•\t" + new_bullet
    assert paras[4].text == "Languages: Python, Airflow, SQL, Docker, Tableau"
    assert paras[4].runs[0].bold is True and paras[4].runs[0].text == "Languages: "


def test_docx_refuses_text_that_needs_more_lines():
    data = make_docx()
    bullet = seg(fileedit.extract("docx", data), "bullet")
    longer = BULLET_1 + " It also emails a summary to every store manager."
    _, results = fileedit.apply("docx", data, {bullet.id: longer})
    assert results[0].status == "too_long"


# ---------------------------------------------------------------- API


@pytest.fixture
def client(monkeypatch):
    for key in ("GEMINI_API_KEY", "GROQ_API_KEY", "SUPABASE_URL", "SUPABASE_ANON_KEY"):
        monkeypatch.setenv(key, "")
    get_settings.cache_clear()
    limits._memory.clear()
    yield TestClient(app_module.app)
    get_settings.cache_clear()


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def test_api_extract_and_render(client):
    data = make_pdf()
    res = client.post("/api/file/extract", json={"filename": "cv.pdf", "file_base64": b64(data)})
    assert res.status_code == 200
    info = res.json()
    assert info["file_type"] == "pdf" and info["editable"] >= 2 and "Jane Doe" in info["text"]

    skills = seg(fileedit.extract("pdf", data), "skills")
    res = client.post(
        "/api/file/render",
        json={
            "file_type": "pdf",
            "file_base64": b64(data),
            "edits": [{"id": skills.id, "text": "SQL, Python, Docker, Airflow, Tableau"}],
        },
    )
    assert res.status_code == 200
    body = res.json()
    assert body["results"][0]["status"] == "applied"
    assert "SQL, Python, Docker" in pdf_text(base64.b64decode(body["file_base64"]))


def test_api_extract_rejects_other_files(client):
    res = client.post("/api/file/extract", json={"filename": "cv.txt", "file_base64": b64(b"hello there " * 20)})
    assert res.status_code == 422


def test_api_file_tailor_writes_ai_edits_and_reports_refusals(client, monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    get_settings.cache_clear()
    data = make_pdf()
    ex = fileedit.extract("pdf", data)
    bullet, skills = seg(ex, "bullet"), seg(ex, "skills")

    async def fake_stream(self, model, system, prompt, schema, temperature):
        if "refused" in system:  # the fix-up round: still too long, so the original stays
            yield Chunk(
                "text",
                json.dumps(
                    {
                        "edits": [
                            {
                                "id": bullet.id,
                                "text": BULLET_1 + " It also emailed every store manager a short daily summary.",
                            }
                        ]
                    }
                ),
            )
            return
        assert "jane@example.com" not in prompt and "Jane Doe" not in prompt  # personal details never sent
        assert f"[EDIT {skills.id}" in prompt
        out = {
            "job": {"title": "Data Engineer", "company": ""},
            "edits": [
                {"id": skills.id, "text": "Python, Airflow, SQL, Docker, Tableau"},
                {"id": bullet.id, "text": BULLET_1 + " It also emailed every store manager a short daily summary."},
            ],
            "keywords": [{"term": "Airflow", "importance": "must"}],
            "project_ideas": [],
            "advice": [],
        }
        yield Chunk("text", json.dumps(out))

    monkeypatch.setattr(app_module.GeminiProvider, "stream", fake_stream)
    res = client.post(
        "/api/file/tailor",
        json={"file_type": "pdf", "file_base64": b64(data), "job_description": JD},
    )
    events = [b for b in res.text.split("\n\n") if b.startswith("event: result")]
    assert events, res.text
    payload = json.loads(events[0].split("data: ", 1)[1])["data"]
    status = {r["id"]: r["status"] for r in payload["results"]}
    assert status == {skills.id: "applied", bullet.id: "too_long"}
    assert payload["edits"] == [{"id": skills.id, "text": "Python, Airflow, SQL, Docker, Tableau"}]
    out = base64.b64decode(payload["file_base64"])
    assert "Languages: Python, Airflow, SQL" in pdf_text(out)
    assert BULLET_1.split(",")[0] in pdf_text(out)
